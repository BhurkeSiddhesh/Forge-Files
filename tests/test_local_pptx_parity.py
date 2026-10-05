"""Work package 17: on-device PDF to PowerPoint (raster-slide mode).

Runs the shipped `ops-pdf-layout.js` + `ops-pdf-pptx.js` (with the vendored PptxGenJS) against the
real vendored pdf.js in headless Chromium. The deck is read back with python-pptx and compared
with the server's `pdf_to_pptx` for the same PDF: slide count, deck size, and every picture's
position and size must match to the EMU; the pictures are compared with an independent PyMuPDF
render of the page. The hidden text layer is checked against PyMuPDF's text positions.

Not compared: how PowerPoint draws the slides (no PowerPoint or LibreOffice here).
"""

from __future__ import annotations

import io
import math
import posixpath
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import numpy as np
import pymupdf as fitz
import pytest
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import local_browser_harness as H  # noqa: E402
import phase4_fixtures as F  # noqa: E402

ROUTE = "/api/pdf/to-pptx"
EMU_PT = 12700

REASON_ENCRYPTED = "this PDF is password protected"
REASON_STRUCTURE = "this file uses features that cannot be processed on this device"
REASON_BIG = "this file exceeds the safe limit for processing on this device"


@pytest.fixture(scope="module")
def lp():
    server, base = H.start_server()
    with H.sync_api.sync_playwright() as pw:
        browser = H.open_browser(pw)
        page = H.LocalPage(browser, base, ["local/ops-pdf-layout.js", "local/ops-pdf-pptx.js"])
        yield page
        browser.close()
    server.shutdown()


def convert(lp, name, data, fields=None, **kw):
    return lp.run(ROUTE, [(name, data, "file", "application/pdf")], fields or {}, **kw)


def deck(result) -> Presentation:
    return Presentation(io.BytesIO(result["bytes"]))


def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == 13]


def text_shapes(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def server_deck(tmp_path: Path, pdf: bytes, dpi: int = 150) -> Presentation:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from scripts.pdf_utils import pdf_to_pptx

    src = tmp_path / "doc.pdf"
    src.write_bytes(pdf)
    out = tmp_path / "srv"
    out.mkdir(exist_ok=True)
    return Presentation(pdf_to_pptx(str(src), str(out), dpi))


def render(pdf: bytes, page: int, dpi: int) -> Image.Image:
    d = fitz.open(stream=pdf, filetype="pdf")
    pix = d[page].get_pixmap(dpi=dpi)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def mean_abs_diff(a: Image.Image, b: Image.Image) -> float:
    w, h = min(a.width, b.width), min(a.height, b.height)
    x = np.asarray(a.convert("RGB").crop((0, 0, w, h))).astype(float)
    y = np.asarray(b.convert("RGB").crop((0, 0, w, h))).astype(float)
    return float(np.abs(x - y).mean())


@pytest.fixture(scope="module")
def article(lp):
    r = convert(lp, "article.pdf", F.article_pdf())
    assert r["status"] == 200, r
    return r


# ── contract ──────────────────────────────────────────────────────────────


def test_local_success_matches_the_server_contract(article):
    assert article["message"] == "PDF converted to PowerPoint"
    assert article["filename"] == "article_forgefiles.org.pptx"
    assert article["body"]["status"] == "success"
    assert article["asked"] == [] and article["api_requests"] == []
    assert article["bytes"][:2] == b"PK"


def test_no_file_is_a_400_with_no_upload(lp):
    r = lp.run(ROUTE, [])
    assert r["status"] == 400 and r["api_requests"] == [] and r["asked"] == []


def test_package_is_a_well_formed_ooxml_zip(article):
    z = zipfile.ZipFile(io.BytesIO(article["bytes"]))
    assert z.testzip() is None
    names = set(z.namelist())
    for part in ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml", "ppt/slideMasters/slideMaster1.xml", "ppt/theme/theme1.xml"):
        assert part in names, part
    for n in names:
        if n.endswith((".xml", ".rels")):
            ET.fromstring(z.read(n))
    types = ET.fromstring(z.read("[Content_Types].xml"))
    overrides = {e.get("PartName") for e in types if e.tag.endswith("Override")}
    defaults = {e.get("Extension") for e in types if e.tag.endswith("Default")}
    for n in names:
        if n == "[Content_Types].xml" or n.endswith("/"):
            continue
        assert ("/" + n) in overrides or n.rsplit(".", 1)[-1] in defaults, n
    for rels in [n for n in names if n.endswith(".rels")]:
        base = Path(rels).parent.parent
        for rel in ET.fromstring(z.read(rels)):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            path = posixpath.normpath(target[1:] if target.startswith("/") else (base / target).as_posix())
            assert path in names, (rels, target)


def test_no_third_party_or_remote_references_are_embedded(article):
    z = zipfile.ZipFile(io.BytesIO(article["bytes"]))
    blob = b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels")))
    assert b"TargetMode=\"External\"" not in blob
    assert b"http://" not in re.sub(rb"http://schemas\.[^\"' <]+|http://purl\.org[^\"' <]+|http://www\.w3\.org[^\"' <]+", b"", blob).replace(b"http://www.", b"")


# ── against the server ────────────────────────────────────────────────────


@pytest.mark.parametrize("make", [F.article_pdf, F.mixed_sizes_pdf, F.scanned_pdf, F.resume_pdf])
def test_same_slides_deck_size_and_picture_boxes_as_the_server(tmp_path, lp, make):
    pdf = make()
    local = deck(convert(lp, "doc.pdf", pdf))
    srv = server_deck(tmp_path, pdf)
    assert len(local.slides) == len(srv.slides)
    assert (local.slide_width, local.slide_height) == (srv.slide_width, srv.slide_height)
    for ls, ss in zip(local.slides, srv.slides):
        (lp_, ), (sp_, ) = pictures(ls), pictures(ss)
        for attr in ("left", "top", "width", "height"):
            assert abs(getattr(lp_, attr) - getattr(sp_, attr)) <= 1, attr  # EMU
        assert lp_.image.size == sp_.image.size  # pixel dimensions


def test_a_mixed_size_pdf_gets_a_deck_sized_to_the_biggest_page_with_pages_centred(lp):
    d = deck(convert(lp, "mixed.pdf", F.mixed_sizes_pdf()))
    assert (d.slide_width, d.slide_height) == (842 * EMU_PT, 792 * EMU_PT)
    boxes = [(p.left, p.top, p.width, p.height) for s in d.slides for p in pictures(s)]
    assert boxes[0] == ((842 - 612) // 2 * EMU_PT, 0, 612 * EMU_PT, 792 * EMU_PT)
    assert boxes[1] == (0, int((792 - 595) / 2 * EMU_PT), 842 * EMU_PT, 595 * EMU_PT)  # 98.5 pt down, 1,250,950 EMU
    assert boxes[2][2:] == (300 * EMU_PT, 200 * EMU_PT)
    for left, top, w, h in boxes:  # nothing sticks out of the slide
        assert left >= 0 and top >= 0 and left + w <= d.slide_width + 1 and top + h <= d.slide_height + 1


@pytest.mark.parametrize("dpi", [72, 150, 300])
def test_dpi_sets_the_picture_pixels(lp, dpi):
    d = deck(convert(lp, "doc.pdf", F.mixed_sizes_pdf(), {"dpi": dpi}))
    sizes = [pictures(s)[0].image.size for s in d.slides]
    px = lambda pt: math.ceil(pt * dpi / 72 - 1e-6)  # the server rounds a partial pixel up, too
    assert sizes == [(px(612), px(792)), (px(842), px(595)), (px(300), px(200))]


@pytest.mark.parametrize("page", [0, 1, 2])
def test_each_picture_is_the_page_not_a_blank_or_cropped_render(lp, page):
    pdf = F.article_pdf()
    d = deck(convert(lp, "doc.pdf", pdf))
    img = Image.open(io.BytesIO(pictures(d.slides[page])[0].image.blob)).convert("RGB")
    ref = render(pdf, page, 150)
    assert img.size == ref.size
    assert mean_abs_diff(img, ref) < 6  # anti-aliasing differs; content, layout and extent do not
    assert np.asarray(img).std() > 20  # not blank


def test_the_pictures_are_lossless_png_like_the_server(article):
    z = zipfile.ZipFile(io.BytesIO(article["bytes"]))
    media = [n for n in z.namelist() if n.startswith("ppt/media/") and not n.endswith("/")]
    assert len(media) == 3 and all(n.endswith(".png") and z.read(n)[:8] == b"\x89PNG\r\n\x1a\n" for n in media)


# ── the hidden text layer ────────────────────────────────────────────────


def test_every_slide_carries_a_fully_transparent_text_layer(article):
    z = zipfile.ZipFile(io.BytesIO(article["bytes"]))
    xml = z.read("ppt/slides/slide1.xml").decode()
    runs = re.findall(r"<a:r>.*?</a:r>", xml, re.S)
    assert len(runs) >= 10
    assert all('<a:alpha val="0"/>' in r for r in runs)  # invisible, but selectable and searchable
    d = Presentation(io.BytesIO(article["bytes"]))
    assert len(text_shapes(d.slides[0])) >= 10


def test_the_text_layer_holds_the_pages_words(article):
    pdf = F.article_pdf()
    d = Presentation(io.BytesIO(article["bytes"]))
    src = fitz.open(stream=pdf, filetype="pdf")
    for i, page in enumerate(src):
        expected = Counter(re.findall(r"[\w'’-]+", page.get_text().lower()))
        got = Counter(re.findall(r"[\w'’-]+", " ".join(s.text_frame.text for s in text_shapes(d.slides[i])).lower()))
        total = sum(expected.values())
        assert sum((expected - got).values()) / total <= 0.005
        assert sum((got - expected).values()) / total <= 0.005


def test_text_boxes_put_the_text_baseline_and_left_edge_where_the_pdf_has_them(article):
    """Each box's baseline and left edge land within 2 pt of PyMuPDF's line origin."""
    pdf = F.article_pdf()
    d = Presentation(io.BytesIO(article["bytes"]))
    src = fitz.open(stream=pdf, filetype="pdf")
    checked = 0
    for i, page in enumerate(src):
        shapes = {s.text_frame.text.strip(): s for s in text_shapes(d.slides[i])}
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                text = "".join(sp["text"] for sp in line["spans"]).strip()
                if not text or text not in shapes:
                    continue
                ox, oy = line["spans"][0]["origin"]
                box = shapes[text]
                size = box.text_frame.paragraphs[0].runs[0].font.size.pt
                assert abs(box.left / EMU_PT - ox) <= 2, (text, box.left / EMU_PT, ox)
                assert abs(box.top / EMU_PT + 0.905 * size - oy) <= 2, (text, box.top / EMU_PT + 0.905 * size, oy)
                checked += 1
    assert checked >= 15


def test_the_text_layer_can_be_turned_off(lp):
    d = deck(convert(lp, "doc.pdf", F.article_pdf(), {"searchable_text": "false"}))
    assert all(not text_shapes(s) for s in d.slides) and all(len(pictures(s)) == 1 for s in d.slides)


def test_pictures_carry_alt_text(article):
    d = Presentation(io.BytesIO(article["bytes"]))
    descr = [pictures(s)[0]._element.xpath(".//p:cNvPr/@descr")[0] for s in d.slides]
    assert descr == ["Page 1 of 3", "Page 2 of 3", "Page 3 of 3"]


def test_a_scanned_pdf_converts_exactly_like_the_server_without_asking(lp):
    r = convert(lp, "scan.pdf", F.scanned_pdf())
    assert r["status"] == 200 and r["asked"] == [] and r["api_requests"] == []
    d = deck(r)
    assert len(d.slides) == 1 and not text_shapes(d.slides[0])  # a picture is the product; nothing to OCR


def test_a_rotated_page_keeps_its_picture_and_skips_only_its_text_layer(lp):
    d = deck(convert(lp, "rot.pdf", F.rotated_page_pdf()))
    assert text_shapes(d.slides[0]) and not text_shapes(d.slides[1])
    w, h = pictures(d.slides[1])[0].image.size
    assert w > h  # page 2 is displayed landscape


# ── validation, as the server does it ────────────────────────────────────


@pytest.mark.parametrize("dpi", [20, 700, 50, 301])
def test_dpi_validation_matches_the_servers_wording(lp, dpi):
    from fastapi.testclient import TestClient

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import main

    r = convert(lp, "doc.pdf", F.plain_pdf(), {"dpi": dpi})
    srv = TestClient(main.app).post(ROUTE, files={"file": ("doc.pdf", F.plain_pdf(), "application/pdf")}, data={"dpi": dpi})
    assert r["status"] == 400 and srv.status_code in (400, 422)
    assert r["detail"].rstrip(".") == str(srv.json()["detail"]).rstrip(".")
    assert r["api_requests"] == [] and r["asked"] == []


def test_too_many_pages_and_too_many_pixels_use_the_servers_messages(lp):
    from fastapi.testclient import TestClient

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import main

    many = F.blank_pdf(201)
    r = convert(lp, "many.pdf", many)
    srv = TestClient(main.app).post(ROUTE, files={"file": ("many.pdf", many, "application/pdf")}, data={"dpi": 150})
    assert r["status"] == 400 and r["detail"] == srv.json()["detail"] == "PDF has too many pages to render at once (max 200)."

    doc = fitz.open()
    doc.new_page(width=5000, height=5000)
    big = doc.tobytes()
    r = convert(lp, "big.pdf", big)
    srv = TestClient(main.app).post(ROUTE, files={"file": ("big.pdf", big, "application/pdf")}, data={"dpi": 150})
    assert r["status"] == 400 and r["detail"] == srv.json()["detail"] == "Page render would exceed 20,000,000 pixels at 150 DPI."
    assert r["api_requests"] == []


# ── the server path is a decision, never an accident ─────────────────────


@pytest.mark.parametrize("make, reason", [
    (F.encrypted_pdf, REASON_ENCRYPTED),
    (lambda: b"%PDF-1.4\nnot really a pdf", REASON_STRUCTURE),
])
def test_unsupported_input_asks_first_and_declining_uploads_nothing(lp, make, reason):
    lp.set_consent(False)
    r = convert(lp, "doc.pdf", make())
    assert r["status"] == 499
    assert [a["reason"] for a in r["asked"]] == [reason]
    assert r["api_requests"] == []


def test_a_supplied_password_asks_first(lp):
    lp.set_consent(False)
    r = convert(lp, "a.pdf", F.plain_pdf(), {"password": "secret"})
    assert [a["reason"] for a in r["asked"]] == [REASON_ENCRYPTED] and r["api_requests"] == []


def test_confirming_sends_exactly_one_request(lp):
    lp.set_consent(True)
    try:
        r = convert(lp, "enc.pdf", F.encrypted_pdf())
    finally:
        lp.set_consent(False)
    assert len(r["asked"]) == 1
    assert len(r["api_requests"]) == 1 and r["api_requests"][0].endswith(ROUTE)


def test_cancelling_stops_without_asking_or_uploading(lp):
    r = convert(lp, "many.pdf", F.many_pages_pdf(60), abort_ms=150)
    assert r.get("threw") == "AbortError"
    assert r["asked"] == [] and r["api_requests"] == []


def test_progress_is_reported_per_page(lp):
    r = convert(lp, "many.pdf", F.many_pages_pdf(4))
    assert r["status"] == 200 and r["progress"][-1] == [4, 4] and len(r["progress"]) == 4
    assert len(deck(r).slides) == 4


def test_mobile_budget_declines_oversized_input_before_parsing(lp):
    lp.page.evaluate("window.matchMedia = () => ({ matches: true })")
    try:
        r = convert(lp, "big.pdf", b"%PDF-1.4\n" + b"0" * (26 * 1024 * 1024))
    finally:
        lp.page.evaluate("delete window.matchMedia")
    assert [a["reason"] for a in r["asked"]] == [REASON_BIG] and r["api_requests"] == []


def test_repeat_conversions_are_independent(lp):
    a = convert(lp, "x.pdf", F.resume_pdf())
    b = convert(lp, "x.pdf", F.resume_pdf())
    assert [pictures(s)[0].image.size for s in deck(a).slides] == [pictures(s)[0].image.size for s in deck(b).slides]
