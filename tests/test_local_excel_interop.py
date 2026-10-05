"""Work packages 32, 33, 34: Interoperability tests between on-device ExcelJS and Python openpyxl.

Verifies:
  1. CSV -> local ExcelJS -> Python openpyxl reads and validates values and formula guards.
  2. XLSX created by Python openpyxl -> local ExcelJS -> Python csv.reader reads and validates values.
  3. Workbooks created by Python openpyxl -> local ExcelJS merge -> Python openpyxl reads merged sheets.
  4. Parity with scripts.excel_utils.
"""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
from openpyxl import Workbook
import pytest

from scripts.excel_utils import csv_to_xlsx, xlsx_to_csv, merge_excel_files

NODE = shutil.which("node")
pytestmark = pytest.mark.skipif(NODE is None, reason="node not available")

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
STATIC_DIR = REPO_ROOT / "public" / "static"

HARNESS_SCRIPT = """
import vm from 'node:vm';
import { readFileSync, writeFileSync } from 'node:fs';
import { join } from 'node:path';

const [, staticDir, apiPath, specPath, outPath] = process.argv;
const spec = JSON.parse(readFileSync(specPath, 'utf8'));

const URLShim = function (...a) { return new URL(...a); };
URLShim.createObjectURL = () => 'blob:mock/1';
URLShim.revokeObjectURL = () => {};

const sandbox = {
    console, Blob, FormData, Response, TextEncoder, TextDecoder,
    Buffer, Uint8Array, ArrayBuffer,
    FileReader: class {
        readAsArrayBuffer(b) { b.arrayBuffer().then(buf => { this.result = buf; this.onload && this.onload(); }); }
        readAsText(b) { b.text().then(txt => { this.result = txt; this.onload && this.onload(); }); }
    },
    Image: class {}, URL: URLShim, setTimeout, clearTimeout, AbortController,
    document: {
        currentScript: null, head: { appendChild() {} },
        createElement: (t) => t === 'canvas' ? { toBlob() {}, getContext: () => null } : {}
    },
    fetch: () => Promise.resolve(new Response('{}'))
};
sandbox.window = sandbox;
sandbox.self = sandbox;
sandbox.window.apiUrl = p => p;
vm.createContext(sandbox);

for (const f of ['vendor/exceljs.min.js', 'vendor/jszip.min.js', 'local/ff-server-gate.js', 'local/ff-local.js', 'local/ops-excel.js']) {
    vm.runInContext(readFileSync(join(staticDir, f), 'utf8'), sandbox, { filename: f });
}

let asked = null;
sandbox.ffConsent.handler = (info) => { asked = info; return false; };

const fd = new FormData();
for (const f of (spec.files || [])) {
    fd.append(f.field, new Blob([readFileSync(f.path)]), f.name || 'file.bin');
}
for (const [k, v] of Object.entries(spec.fields || {})) {
    fd.append(k, String(v));
}

const res = await sandbox.ffProcess(apiPath, fd);
const body = await res.json();
const out = { status: res.status, detail: body.detail, message: body.message, filename: body.filename, asked };
if (res.ok) {
    const item = sandbox.ffLocal.resolve(body.download_token);
    writeFileSync(outPath, Buffer.from(await item.blob.arrayBuffer()));
}
console.log(JSON.stringify(out));
"""


def _run_local_excel(tmp_path: Path, api_path: str, files: list[dict], fields: dict | None = None) -> tuple[dict, Path]:
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps({"files": files, "fields": fields or {}}), encoding="utf-8")
    ext = ".xlsx" if "xlsx" in api_path or "merge" in api_path else ".csv"
    out_path = tmp_path / f"out_excel_result{ext}"
    if out_path.exists():
        out_path.unlink()

    proc = subprocess.run(
        [NODE, "--input-type=module", "-e", HARNESS_SCRIPT, str(STATIC_DIR), api_path, str(spec_path), str(out_path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )
    assert proc.returncode == 0, f"Node script error: {proc.stderr}\n{proc.stdout}"
    output_line = proc.stdout.strip().splitlines()[-1]
    return json.loads(output_line), out_path


def test_interop_csv_to_xlsx(tmp_path: Path):
    """CSV converted by local ExcelJS can be loaded and read cleanly by openpyxl."""
    csv_file = tmp_path / "products.csv"
    csv_content = (
        "Item,Price,Qty,Formula_Test\n"
        "Widget,19.99,10,=2+2\n"
        "Gadget,5.50,100,+10-5\n"
        "Doodad,0.99,500,@SUM(1)\n"
    )
    csv_file.write_text(csv_content, encoding="utf-8")

    result, out_xlsx = _run_local_excel(
        tmp_path,
        "/api/excel/csv-to-xlsx",
        [{"field": "file", "path": str(csv_file), "name": "products.csv"}],
    )
    assert result["status"] == 200
    assert result["message"] == "CSV converted to XLSX"
    assert result["filename"] == "products_forgefiles.org.xlsx"

    # Verify with openpyxl
    wb = openpyxl.load_workbook(out_xlsx)
    assert len(wb.sheetnames) == 1
    ws = wb.active
    assert ws.title == "products"

    # Headers
    assert ws["A1"].value == "Item"
    assert ws["B1"].value == "Price"
    assert ws["D1"].value == "Formula_Test"

    # Numeric rows
    assert ws["A2"].value == "Widget"
    assert ws["B2"].value == 19.99
    assert ws["C2"].value == 10

    # Formula injection guard: values starting with = + @ are preserved as strings and given '@' text format
    assert ws["D2"].value == "=2+2"
    assert ws["D2"].number_format == "@"

    assert ws["D3"].value == "+10-5"
    assert ws["D3"].number_format == "@"

    assert ws["D4"].value == "@SUM(1)"
    assert ws["D4"].number_format == "@"


def test_interop_xlsx_to_csv(tmp_path: Path):
    """XLSX generated by openpyxl is cleanly exported to CSV by local ExcelJS."""
    xlsx_file = tmp_path / "inventory.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["SKU", "Description", "Stock"])
    ws.append(["A101", "Red, Shiny Apple", 50])
    ws.append(["B202", 'Banana ("Cavendish")', 120])
    wb.save(xlsx_file)

    result, out_csv = _run_local_excel(
        tmp_path,
        "/api/excel/xlsx-to-csv",
        [{"field": "file", "path": str(xlsx_file), "name": "inventory.xlsx"}],
    )
    assert result["status"] == 200
    assert result["filename"] == "inventory_forgefiles.org.csv"

    # Read back with Python csv.reader
    with out_csv.open("r", encoding="utf-8", newline="") as f:
        reader = list(csv.reader(f))

    assert len(reader) == 3
    assert reader[0] == ["SKU", "Description", "Stock"]
    assert reader[1] == ["A101", "Red, Shiny Apple", "50"]
    assert reader[2] == ["B202", 'Banana ("Cavendish")', "120"]


def test_interop_merge_excel(tmp_path: Path):
    """Multiple XLSX workbooks generated by openpyxl are cleanly merged by local ExcelJS."""
    wb1 = Workbook()
    ws1 = wb1.active
    ws1.title = "Sales"
    ws1.append(["Q1", 1000])
    f1 = tmp_path / "DeptA.xlsx"
    wb1.save(f1)

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Sales"
    ws2.append(["Q2", 2000])
    f2 = tmp_path / "DeptB.xlsx"
    wb2.save(f2)

    result, out_merged = _run_local_excel(
        tmp_path,
        "/api/excel/merge",
        [
            {"field": "files", "path": str(f1), "name": "DeptA.xlsx"},
            {"field": "files", "path": str(f2), "name": "DeptB.xlsx"},
        ],
    )
    assert result["status"] == 200
    assert result["message"] == "Excel files merged"
    assert result["filename"].startswith("merged_")

    merged_wb = openpyxl.load_workbook(out_merged)
    assert len(merged_wb.sheetnames) == 2
    assert "DeptA_Sales" in merged_wb.sheetnames
    assert "DeptB_Sales" in merged_wb.sheetnames

    assert merged_wb["DeptA_Sales"]["A1"].value == "Q1"
    assert merged_wb["DeptA_Sales"]["B1"].value == 1000

    assert merged_wb["DeptB_Sales"]["A1"].value == "Q2"
    assert merged_wb["DeptB_Sales"]["B1"].value == 2000


def test_parity_csv_to_xlsx_with_server(tmp_path: Path):
    """Compare local ExcelJS csv-to-xlsx with scripts.excel_utils.csv_to_xlsx."""
    csv_file = tmp_path / "compare.csv"
    csv_file.write_text("Col1,Col2\nVal1,123\nVal2,456\n", encoding="utf-8")

    # Local output
    result, out_local = _run_local_excel(
        tmp_path,
        "/api/excel/csv-to-xlsx",
        [{"field": "file", "path": str(csv_file), "name": "compare.csv"}],
    )
    assert result["status"] == 200

    # Server output
    server_dir = tmp_path / "server_out"
    server_dir.mkdir(exist_ok=True)
    out_server = Path(csv_to_xlsx(str(csv_file), str(server_dir)))

    wb_local = openpyxl.load_workbook(out_local)
    wb_server = openpyxl.load_workbook(out_server)

    assert wb_local.sheetnames == wb_server.sheetnames
    ws_local = wb_local.active
    ws_server = wb_server.active

    local_rows = list(ws_local.iter_rows(values_only=True))
    server_rows = list(ws_server.iter_rows(values_only=True))

    assert len(local_rows) == len(server_rows)
    assert local_rows[0] == ("Col1", "Col2")
    assert local_rows[1][0] == server_rows[1][0]
    assert local_rows[2][0] == server_rows[2][0]
