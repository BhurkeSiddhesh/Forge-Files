"""Work package 42: on-device PDF to Word.

Runs the shipped `ops-pdf-layout.js` + `ops-pdf-word.js` (with the vendored docx writer) against
the real vendored pdf.js in headless Chromium. The .docx is read back with python-docx and its
XML, never trusted because a download appeared. Text accuracy is measured against PyMuPDF's own
reading of the source and against what the server's pdf2docx produces for the same PDF.

Pagination was NOT compared against a rendered copy (no Word or LibreOffice here); these tests
prove structure (one section per source page, exact line spacing, tab stops, tables, lists), not
how Word lays the pages out.
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import docx
import pymupdf as fitz
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
import local_browser_harness as H  # noqa: E402
import phase4_fixtures as F  # noqa: E402

ROUTE = "/api/pdf/convert-to-word"
STREAM = "/api/pdf/convert-to-word-stream"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

REASON_SCAN = "some pages are scans that need text recognition (OCR)"
REASON_STRUCTURE = "this file uses features that cannot be processed on this device"
REASON_ENCRYPTED = "this PDF is password protected"
REASON_AI = "AI layout analysis (an option you chose) runs on the server"
REASON_BIG = "this file exceeds the safe limit for processing on this device"


@pytest.fixture(scope="module")
def lp():
    server, base = H.start_server()
    with H.sync_api.sync_playwright() as pw:
        browser = H.open_browser(pw)
        page = H.LocalPage(browser, base, ["local/ops-pdf-layout.js", "local/ops-pdf-word.js"])
        yield page
        browser.close()
    server.shutdown()


def convert(lp, name, data, route=ROUTE, **kw):
    return lp.run(route, [(name, data, "file", "application/pdf")], **kw)


def doc_of(result) -> docx.document.Document:
    return docx.Document(io.BytesIO(result["bytes"]))


def body_items(d):
    """Body children as ('p', Paragraph) / ('t', Table) in document order."""
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            out.append(("p", docx.text.paragraph.Paragraph(el, d)))
        elif tag == "tbl":
            out.append(("t", docx.table.Table(el, d)))
    return out


def paragraphs(d):
    return [p for kind, p in body_items(d) if kind == "p"]


def text_of(d) -> str:
    parts = []
    for kind, item in body_items(d):
        if kind == "p":
            parts.append(item.text)
        else:
            for row in item.rows:
                parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def words(text: str) -> Counter:
    return Counter(re.findall(r"[\w'’-]+", text.lower()))


def source_words(pdf: bytes) -> Counter:
    d = fitz.open(stream=pdf, filetype="pdf")
    return words(" ".join(p.get_text() for p in d))


def recall_precision(expected: Counter, got: Counter):
    total = sum(expected.values())
    return 1 - sum((expected - got).values()) / total, 1 - sum((got - expected).values()) / sum(got.values())


@pytest.fixture(scope="module")
def resume(lp):
    r = convert(lp, "resume.pdf", F.resume_pdf())
    assert r["status"] == 200, r
    return r


@pytest.fixture(scope="module")
def article(lp):
    r = convert(lp, "article.pdf", F.article_pdf())
    assert r["status"] == 200, r
    return r


# ── contract ──────────────────────────────────────────────────────────────


def test_local_success_matches_the_server_contract(resume):
    assert resume["message"] == "Converted to Word (Standard)"
    assert resume["filename"] == "resume_forgefiles.org.docx"
    assert resume["asked"] == [] and resume["api_requests"] == []
    assert resume["bytes"][:2] == b"PK"


def test_the_stream_route_the_ui_uses_returns_the_same_document(lp, resume):
    r = convert(lp, "resume.pdf", F.resume_pdf(), route=STREAM)
    assert r["status"] == 200 and r["contentType"].startswith("text/event-stream")
    assert [e["event"] for e in r["events"]] == ["start", "complete"]
    done = r["events"][1]
    assert done["message"] == "Converted to Word (Standard)" and done["method"] == "standard"
    assert done["filename"] == "resume_forgefiles.org.docx" and done["download_token"].startswith("ffLocal:")
    assert r["api_requests"] == []
    assert text_of(doc_of(r)) == text_of(doc_of(resume))


def test_no_file_is_a_400_with_no_upload(lp):
    r = lp.run(ROUTE, [])
    assert r["status"] == 400 and r["api_requests"] == [] and r["asked"] == []


# ── package ───────────────────────────────────────────────────────────────


def test_package_is_a_well_formed_ooxml_zip(resume):
    z = zipfile.ZipFile(io.BytesIO(resume["bytes"]))
    assert z.testzip() is None
    names = set(z.namelist())
    for part in ("[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml", "word/numbering.xml"):
        assert part in names, part
    for n in names:
        if n.endswith((".xml", ".rels")):
            ET.fromstring(z.read(n))  # well-formed everywhere
    types = ET.fromstring(z.read("[Content_Types].xml"))
    overrides = {e.get("PartName") for e in types if e.tag.endswith("Override")}
    defaults = {e.get("Extension") for e in types if e.tag.endswith("Default")}
    for n in names:
        if n == "[Content_Types].xml" or n.endswith("/"):
            continue
        assert ("/" + n) in overrides or n.rsplit(".", 1)[-1] in defaults, n
    # every relationship target exists
    for rels in [n for n in names if n.endswith(".rels")]:
        base = Path(rels).parent.parent
        for rel in ET.fromstring(z.read(rels)):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            path = (Path(target[1:]) if target.startswith("/") else base / target).as_posix()
            assert path in names, (rels, target)


def test_word_opens_it_one_section_per_page_sized_like_the_source(article):
    d = doc_of(article)
    assert len(d.sections) == 3
    for s in d.sections:
        assert round(s.page_width.pt) == 595 and round(s.page_height.pt) == 842
    # every section but the first starts on a new page
    starts = [s.start_type for s in d.sections]
    assert all(str(t).startswith("NEW_PAGE") for t in starts[1:])


# ── structure and styling ────────────────────────────────────────────────


def test_headings_runs_fonts_and_sizes(resume):
    d = doc_of(resume)
    ps = paragraphs(d)
    assert ps[0].text == "Jane Q. Example" and ps[0].style.name == "Heading 1"
    assert ps[0].runs[0].bold and ps[0].runs[0].font.size.pt == 24
    assert ps[2].text == "EXPERIENCE" and ps[2].style.name == "Heading 2"
    company = ps[3]
    assert company.runs[0].text == "Acme Corporation, Staff Engineer" and company.runs[0].bold
    assert company.runs[0].font.size.pt == 11
    assert company.runs[0].font.name == "Arial"  # Liberation Sans maps to its metric twin


def test_a_right_aligned_date_is_a_right_tab_not_spaces(resume):
    d = doc_of(resume)
    line = next(p for p in paragraphs(d) if p.text.startswith("Acme Corporation"))
    assert line.text == "Acme Corporation, Staff Engineer\t2021 – 2025"
    stops = [(t.alignment, round(t.position.pt)) for t in line.paragraph_format.tab_stops]
    page = d.sections[0]
    text_right = page.page_width.pt - page.right_margin.pt - page.left_margin.pt
    assert stops == [(WD_TAB_ALIGNMENT.RIGHT, round(text_right))]


def test_bullets_are_real_word_list_items(resume):
    d = doc_of(resume)
    items = [p for p in paragraphs(d) if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert [p.text for p in items] == ["Led a team of eight engineers on the billing platform",
                                       "Cut p99 latency by 40% through caching and batching",
                                       "Mentored four engineers into senior roles"]
    assert all("•" not in p.text for p in items)  # the glyph comes from the list, not the text
    assert all(p.paragraph_format.left_indent.pt > 18 for p in items)


def test_line_spacing_is_exact_so_pagination_does_not_drift(resume):
    d = doc_of(resume)
    for p in paragraphs(d):
        spacing = p._p.pPr.find(W + "spacing")
        assert spacing is not None
        assert spacing.get(W + "lineRule") == "exact" and int(spacing.get(W + "line")) >= 200


def test_wrapped_text_is_one_paragraph_that_keeps_the_pdfs_wrap_width(lp):
    d = doc_of(convert(lp, "narrow.pdf", F.narrow_block_pdf()))
    wide, narrow = paragraphs(d)
    assert wide.text.startswith("The committee reviewed") and wide.text.endswith("takes place.") and "\n" not in wide.text
    assert narrow.text.startswith("Subsequent analysis") and narrow.text.endswith("reliability work.") and "\n" not in narrow.text
    # The narrow block is indented from the right by the difference, so Word wraps where the PDF did.
    assert (wide.paragraph_format.right_indent or 0) < 4 * 12700
    assert 150 < narrow.paragraph_format.right_indent.pt < 230


def test_hyperlinks_are_external_relationships_and_only_safe_ones(lp, article):
    d = doc_of(article)
    targets = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert targets == ["https://example.com/more"]
    r = convert(lp, "links.pdf", F.hostile_link_pdf())
    d2 = doc_of(r)
    assert [x.target_ref for x in d2.part.rels.values() if x.reltype.endswith("/hyperlink")] == ["https://example.org/ok"]
    xml = zipfile.ZipFile(io.BytesIO(r["bytes"])).read("word/_rels/document.xml.rels").decode()
    assert "javascript:" not in xml and "file:" not in xml


def test_an_embedded_image_is_kept_inline_and_decodes(article):
    d = doc_of(article)
    assert len(d.inline_shapes) == 1
    shape = d.inline_shapes[0]
    assert round(shape.width.pt) == 240 and round(shape.height.pt) == 120
    blob = next(r.target_part.blob for r in d.part.rels.values() if r.reltype.endswith("/image"))
    img = Image.open(io.BytesIO(blob))
    assert img.convert("RGB").getpixel((60, 60)) == (200, 40, 40)


def test_unicode_text_survives(lp):
    r = convert(lp, "uni.pdf", F.unicode_pdf())
    text = text_of(doc_of(r))
    assert "Café crème – naïve façade" in text
    assert "Привет мир" in text and "Γειά" in text


# ── text accuracy ─────────────────────────────────────────────────────────


@pytest.mark.parametrize("make", [F.resume_pdf, F.article_pdf, F.plain_pdf, F.unicode_pdf])
def test_every_source_word_is_kept(lp, make):
    pdf = make()
    r = convert(lp, "doc.pdf", pdf)
    recall, precision = recall_precision(source_words(pdf), words(text_of(doc_of(r))))
    assert recall >= 0.995 and precision >= 0.995, (recall, precision)


def test_text_is_at_least_as_accurate_as_the_servers_pdf2docx(tmp_path, lp):
    pytest.importorskip("pdf2docx")
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from scripts.pdf_utils import pdf_to_docx

    pdf = F.resume_pdf()
    src = tmp_path / "resume.pdf"
    src.write_bytes(pdf)
    out_dir = tmp_path / "srv"
    out_dir.mkdir()
    server_doc = docx.Document(pdf_to_docx(str(src), str(out_dir)))
    expected = source_words(pdf)
    server_recall, _ = recall_precision(expected, words(text_of(server_doc)))
    local_recall, _ = recall_precision(expected, words(text_of(doc_of(convert(lp, "resume.pdf", pdf)))))
    assert local_recall >= 0.995
    assert local_recall >= server_recall - 0.001, (local_recall, server_recall)


# ── tables ────────────────────────────────────────────────────────────────

GRID = [list(r) for r in F.ROWS]


def table_of(r):
    d = doc_of(r)
    tables = d.tables
    assert len(tables) == 1, len(tables)
    return d, tables[0]


def test_a_ruled_table_becomes_a_real_word_table(lp):
    r = convert(lp, "t.pdf", F.ruled_table_pdf())
    _, t = table_of(r)
    assert [[c.text for c in row.cells] for row in t.rows] == GRID
    borders = t._tbl.xpath(".//w:tcBorders/w:top")
    assert borders and all(b.get(W + "val") == "single" for b in borders)


def test_numbers_in_a_ruled_table_stay_right_aligned(lp):
    r = convert(lp, "t.pdf", F.ruled_table_pdf())
    _, t = table_of(r)
    for row in t.rows[1:]:
        assert row.cells[0].paragraphs[0].alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)
        assert all(row.cells[c].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT for c in (1, 2, 3))
    assert all(run.bold for run in t.rows[0].cells[0].paragraphs[0].runs)  # header row stays bold


def test_a_borderless_table_is_found_from_alignment_alone(lp):
    r = convert(lp, "t.pdf", F.borderless_table_pdf())
    d, t = table_of(r)
    assert [[c.text for c in row.cells] for row in t.rows] == GRID
    assert not t._tbl.xpath(".//w:tcBorders/w:top[@w:val='single']")  # no lines were invented
    assert paragraphs(d)[0].text == "Order lines"


def test_a_table_between_two_paragraphs_keeps_the_order(lp):
    r = convert(lp, "t.pdf", F.ruled_table_pdf(with_paragraphs=True))
    d = doc_of(r)
    kinds = [k for k, _ in body_items(d)]
    first_t = kinds.index("t")
    assert kinds[0] == "p" and first_t > 0 and "p" in kinds[first_t + 1:]
    texts = [i.text for k, i in body_items(d) if k == "p"]
    assert texts[0].startswith("Quarterly order summary") and texts[-1].startswith("Subsequent analysis")


def test_a_table_with_merged_cells_asks_instead_of_guessing(lp):
    lp.set_consent(False)
    r = convert(lp, "m.pdf", F.merged_table_pdf())
    assert r["status"] == 499 and [a["reason"] for a in r["asked"]] == [REASON_STRUCTURE] and r["api_requests"] == []


# ── the server path is a decision, never an accident ─────────────────────


@pytest.mark.parametrize("make, reason", [
    (F.scanned_pdf, REASON_SCAN),
    (F.two_column_pdf, REASON_STRUCTURE),
    (F.three_column_pdf, REASON_STRUCTURE),
    (F.encrypted_pdf, REASON_ENCRYPTED),
    (lambda: b"%PDF-1.4\nnot really a pdf", REASON_STRUCTURE),
])
@pytest.mark.parametrize("route", [ROUTE, STREAM])
def test_unsupported_input_asks_first_and_declining_uploads_nothing(lp, make, reason, route):
    lp.set_consent(False)
    r = convert(lp, "doc.pdf", make(), route=route)
    assert r["status"] == 499
    assert [a["reason"] for a in r["asked"]] == [reason]
    assert r["api_requests"] == []


def test_ai_mode_is_a_server_feature_and_asks_first(lp):
    lp.set_consent(False)
    r = lp.run(STREAM, [("a.pdf", F.plain_pdf(), "file", "application/pdf")], {"use_ai": "true"})
    assert [a["reason"] for a in r["asked"]] == [REASON_AI] and r["api_requests"] == []


def test_a_supplied_password_asks_first(lp):
    lp.set_consent(False)
    r = lp.run(STREAM, [("a.pdf", F.plain_pdf(), "file", "application/pdf")], {"password": "secret"})
    assert [a["reason"] for a in r["asked"]] == [REASON_ENCRYPTED] and r["api_requests"] == []


@pytest.mark.parametrize("route", [ROUTE, STREAM])
def test_confirming_sends_exactly_one_request(lp, route):
    lp.set_consent(True)
    try:
        r = convert(lp, "scan.pdf", F.scanned_pdf(), route=route)
    finally:
        lp.set_consent(False)
    assert len(r["asked"]) == 1
    assert len(r["api_requests"]) == 1 and r["api_requests"][0].endswith(route)


def test_cancelling_stops_without_asking_or_uploading(lp):
    r = convert(lp, "many.pdf", F.many_pages_pdf(80), abort_ms=60, route=STREAM)
    assert r.get("threw") == "AbortError"
    assert r["asked"] == [] and r["api_requests"] == []


def test_progress_is_reported_per_page(lp):
    r = convert(lp, "many.pdf", F.many_pages_pdf(5))
    assert r["status"] == 200 and r["progress"][-1] == [5, 5]
    assert len(doc_of(r).sections) == 5


def test_mobile_budget_declines_oversized_input_before_parsing(lp):
    lp.page.evaluate("window.matchMedia = () => ({ matches: true })")
    try:
        r = convert(lp, "big.pdf", b"%PDF-1.4\n" + b"0" * (26 * 1024 * 1024))
    finally:
        lp.page.evaluate("delete window.matchMedia")
    assert [a["reason"] for a in r["asked"]] == [REASON_BIG] and r["api_requests"] == []


def test_repeat_conversions_are_independent(lp):
    a = convert(lp, "resume.pdf", F.resume_pdf())
    b = convert(lp, "resume.pdf", F.resume_pdf())
    assert text_of(doc_of(a)) == text_of(doc_of(b))


def test_a_rotated_page_asks_instead_of_producing_a_wrongly_oriented_page(lp):
    lp.set_consent(False)
    r = convert(lp, "rot.pdf", F.rotated_page_pdf())
    assert r["status"] == 499 and [a["reason"] for a in r["asked"]] == [REASON_STRUCTURE] and r["api_requests"] == []
